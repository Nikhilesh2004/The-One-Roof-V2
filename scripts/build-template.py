"""
The bulk-upload spreadsheet, in one place.

    python scripts/build-template.py
        Rewrites the blank template in docs/import/ (xlsx + csv), with two
        example rows. This is what the photo team fills in.

    python scripts/build-template.py <products.json> <out.xlsx>
        Writes the same layout filled with real rows — used for the export of
        the live catalogue (see scripts/export-catalogue.ts).

Both come from the same column list, so a filled export always imports
through the same checks as a sheet the team typed by hand.

Column order is also the importer's contract: src/components/admin/BulkUpload
reads these headers by name, not by position, but keeping one list here is
what stops the two drifting apart.
"""
import csv
import io
import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

PHOTOS = [f"photo_{i}" for i in range(1, 7)]
HEAD = (
    ["title"] + PHOTOS + ["thumbnail", "price", "mrp", "stock", "category",
    "subCategory", "description", "occasions", "countryOfOrigin", "manufacturer",
    "genericName", "netQuantity", "dimensions", "weight", "finish",
    "monthYearOfImport", "status", "featured", "shortCaption", "shortSticker",
    # Last on purpose: the team leaves it blank. It only matters for products
    # that already have a web address people have shared.
    "slug"]
)
REQUIRED = {"title", "photo_1", "price", "stock", "category", "status"}

WIDTH = {"title": 36, "thumbnail": 30, "description": 48, "subCategory": 20,
         "occasions": 22, "manufacturer": 24, "countryOfOrigin": 16,
         "genericName": 18, "netQuantity": 14, "dimensions": 16, "weight": 12,
         "finish": 20, "monthYearOfImport": 18, "shortCaption": 28,
         "shortSticker": 14, "slug": 40}
for _p in PHOTOS:
    WIDTH[_p] = 34

SECTIONS = [("decor", "Home & Decor"), ("pooja", "Pooja Items"),
            ("fashion", "Jewellery & Fashion"), ("bags", "Bags"),
            ("shoes", "Shoes"), ("watches", "Watches"),
            ("gifting", "Gifting & Dining"), ("lifestyle", "Tech & Lifestyle")]

EXAMPLES = [
    {"title": "Krishna with flute - Venugopala form",
     "photo_1": "krishna-flute-01.jpg", "photo_2": "krishna-flute-02.jpg",
     "photo_3": "krishna-flute-03.jpg", "thumbnail": "krishna-flute-thumb.jpg",
     "price": 3950, "mrp": 4500, "stock": 4, "category": "pooja",
     "subCategory": "Brass Idols",
     "description": "Brass Krishna standing with a flute, on a lotus base.",
     "occasions": "festival;wedding", "countryOfOrigin": "India",
     "manufacturer": "Sri Balaji Handicrafts", "genericName": "Brass idol",
     "netQuantity": "1 piece", "dimensions": "9 x 4 x 3 in", "weight": "850 g",
     "finish": "Antique brass", "monthYearOfImport": "09/2026", "status": "live",
     "featured": "yes", "shortCaption": "The one everyone asks for",
     "shortSticker": "Bestseller"},
    {"title": "Leather sling bag - tan",
     "photo_1": "tan-sling-01.jpg", "photo_2": "tan-sling-02.jpg",
     "price": 1299, "mrp": 1799, "stock": 12, "category": "bags",
     "subCategory": "Slings & Crossbody",
     "description": "Compact tan sling with an adjustable strap and two zip pockets.",
     "occasions": "birthday;everyday", "countryOfOrigin": "India",
     "manufacturer": "Hidesign Leathers", "genericName": "Sling bag",
     "netQuantity": "1 piece", "dimensions": "10 x 7 x 3 in", "weight": "420 g",
     "finish": "Genuine leather", "monthYearOfImport": "08/2026",
     "status": "live", "featured": "no"},
]

GUIDE = [
    ("The One Roof - bulk product upload", True),
    ("", False),
    ("One row is one product. Fill in the Products sheet, put ALL the images in ONE folder.", False),
    ("Then in the CMS, open Bulk upload, drop in this file and the images, and check the preview.", False),
    ("", False),
    ("Gold headers are required. Green is the thumbnail. Grey is optional - leave blank if you don't have it.", False),
    ("", False),
    ("PHOTOS - one per column", True),
    ("Put each photo's filename in its own column: photo_1, photo_2, and so on. Up to six.", False),
    ("Leave the rest blank. Most products need two or three.", False),
    ("", False),
    ("      photo_1   krishna-flute-01.jpg", False),
    ("      photo_2   krishna-flute-02.jpg", False),
    ("      photo_3   krishna-flute-03.jpg", False),
    ("", False),
    ("These are the pictures on the product's own page, in the order you list them.", False),
    ("", False),
    ("THUMBNAIL - the picture in the shop grid", True),
    ("A SEPARATE image, not one of the photos above. It is the one picture a customer sees", False),
    ("while browsing, before they click in - so it is the one that sells the product.", False),
    ("", False),
    ("      thumbnail   krishna-flute-thumb.jpg", False),
    ("", False),
    ("      - It goes in the same folder as the photos.", False),
    ("      - Name it after the product with -thumb on the end.", False),
    ("      - Roughly square works best. The grid crops to a square.", False),
    ("      - LEAVE IT BLANK and photo_1 is used instead. That is fine.", False),
    ("", False),
    ("FILENAME RULES - the part nobody can fix afterwards", True),
    ("      - No spaces. Use hyphens.", False),
    ("      - Every filename in the sheet must be in the folder, spelled exactly the same.", False),
    ("        Krishna-Flute-01.JPG and krishna-flute-01.jpg are different files.", False),
    ("      - Send the images that come out of the ChatGPT step, not the originals off the phone.", False),
    ("", False),
    ("THE FIDDLY COLUMNS", True),
    ("price, mrp, stock   Numbers only. 3950, never Rs 3,950. mrp is the struck-through price.", False),
    ("category            Pick from the dropdown. Codes are on the 'Section codes' sheet.", False),
    ("subCategory         Your own grouping, e.g. 'Brass Idols'. Spell it the same way every time.", False),
    ("occasions           Any of: birthday, wedding, festival, everyday - semicolons between them.", False),
    ("status              live = on the website now. draft = ready but held back. hidden = retired.", False),
    ("featured            yes puts it in the Shoppable Shorts rail. A handful only.", False),
    ("slug                LEAVE BLANK. It is the product's web address, filled in automatically.", False),
    ("                    Only filled for products that already exist, so their links keep working.", False),
    ("", False),
    ("UPLOADING THE SAME SHEET TWICE IS SAFE", True),
    ("A product that already exists is updated, not copied. So you can fix a price in the sheet", False),
    ("and upload it again.", False),
    ("", False),
    ("SEND THE FIRST FIFTY FIRST.", True),
    ("A mistake caught on fifty rows costs an hour. Caught on a thousand, it costs a week.", False),
]


def build(path, rows, example=False):
    wb = Workbook()
    ws = wb.active
    ws.title = "Products"

    gold = PatternFill("solid", fgColor="C8A24A")
    grey = PatternFill("solid", fgColor="EDEDED")
    green = PatternFill("solid", fgColor="BFD9D2")
    line = Border(bottom=Side(style="thin", color="BBBBBB"))

    for i, h in enumerate(HEAD, 1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = Font(bold=True, size=11, color="1A1A1A")
        c.fill = gold if h in REQUIRED else (green if h == "thumbnail" else grey)
        c.alignment = Alignment(horizontal="left", vertical="center")
        c.border = line
        ws.column_dimensions[get_column_letter(i)].width = WIDTH.get(h, 14)
    ws.row_dimensions[1].height = 24

    for r, row in enumerate(rows, 2):
        for i, h in enumerate(HEAD, 1):
            v = row.get(h, "")
            cell = ws.cell(row=r, column=i, value=v if v != "" else None)
            if example:
                cell.font = Font(italic=True, color="777777")
            cell.alignment = Alignment(vertical="top", wrap_text=(h == "description"))

    ws.freeze_panes = "B2"

    def dropdown(values, col, prompt):
        d = DataValidation(type="list", formula1='"' + ",".join(values) + '"',
                           allow_blank=True, showDropDown=False)
        d.promptTitle, d.prompt, d.showInputMessage = "Pick from the list", prompt, True
        d.errorTitle, d.error, d.showErrorMessage = "Not a valid value", prompt, True
        ws.add_data_validation(d)
        L = get_column_letter(HEAD.index(col) + 1)
        d.add(f"{L}2:{L}5000")

    dropdown([s for s, _ in SECTIONS], "category",
             "One of the eight section codes - see the 'Section codes' sheet.")
    dropdown(["live", "draft", "hidden"], "status",
             "live = on the website. draft = held back. hidden = retired.")
    dropdown(["yes", "no"], "featured",
             "yes puts it in the Shoppable Shorts rail. Keep this to a handful.")

    s2 = wb.create_sheet("Section codes")
    s2.append(["Code (use this)", "Section shown on the website"])
    for a, b in SECTIONS:
        s2.append([a, b])
    for c in s2[1]:
        c.font = Font(bold=True)
        c.fill = gold
    s2.column_dimensions["A"].width = 18
    s2.column_dimensions["B"].width = 30

    s3 = wb.create_sheet("How to fill this in")
    for i, (t, bold) in enumerate(GUIDE, 1):
        c = s3.cell(row=i, column=1, value=t)
        c.font = Font(bold=bold, size=12 if bold else 11,
                      color="8A6D1F" if bold else "1A1A1A")
    s3.column_dimensions["A"].width = 112
    s3.sheet_view.showGridLines = False

    wb.save(path)


def write_csv(path, rows):
    with io.open(path, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(HEAD)
        for row in rows:
            w.writerow([row.get(h, "") for h in HEAD])


if __name__ == "__main__":
    if len(sys.argv) == 3:
        rows = json.load(io.open(sys.argv[1], encoding="utf-8"))
        build(sys.argv[2], rows)
        print(f"{len(rows)} products -> {sys.argv[2]}")
    else:
        build("docs/import/products-template.xlsx", EXAMPLES, example=True)
        write_csv("docs/import/products-template.csv", EXAMPLES)
        # Served by the site, so Bulk upload's "Download the template" always
        # hands out the same file this script just wrote.
        build("public/bulk-upload-template.xlsx", EXAMPLES, example=True)
        print("blank template -> docs/import/ and public/bulk-upload-template.xlsx")
